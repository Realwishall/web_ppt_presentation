#!/usr/bin/env python3
"""
Phase 1 — deterministic extraction.  No LLM involved.

    python tools/extract.py "<path to .pptx>" [--out build] [--dpi 150]
    python tools/extract.py "<path to chapter folder>" --all

Produces build/<deck-slug>/
    slide-001.png …          rendered at --dpi (default 150), one per slide
    slides.json              per-slide text runs, tables as real grids, notes
    media/                   images embedded in the pptx, native resolution
    manifest.json            authoritative slide count + source metadata

Two independent views of every slide — pixels (catches equations and diagrams
that live inside shapes) and XML (catches exact wording and true table
dimensions) — so the render pass never has to guess.

Rendering: on Windows with PowerPoint installed it exports the slides natively
(highest fidelity, no extra tooling — needs `pip install pywin32`). Otherwise it
falls back to LibreOffice + poppler. Force either with --renderer.

Requires: python-pptx, plus EITHER PowerPoint + pywin32 OR LibreOffice (soffice)
and poppler-utils (pdftoppm).
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("python-pptx is required:  pip install python-pptx")


# ----------------------------------------------------------------- helpers --

def slugify(name: str) -> str:
    s = re.sub(r"\.pptx$", "", name, flags=re.I)
    s = re.sub(r"[^\w\s\-.]", "", s)
    s = re.sub(r"[\s_]+", "-", s.strip())
    s = re.sub(r"-{2,}", "-", s)
    return s.strip("-.").lower()


def which_or_die(*names: str) -> str:
    for n in names:
        p = shutil.which(n)
        if p:
            return p
    sys.exit(f"required executable not found on PATH: {' / '.join(names)}")


def is_lock_file(p: Path) -> bool:
    """Office leaves ~$foo.pptx lock files around — never a real deck."""
    return p.name.startswith("~$")


# ------------------------------------------------------------- rendering --

def existing_pngs(out_dir: Path) -> list[Path]:
    return sorted(out_dir.glob("slide-*.png"),
                  key=lambda p: int(re.search(r"(\d+)", p.stem).group(1)))


def normalise_names(out_dir: Path) -> list[Path]:
    """Renderers disagree about zero-padding (slide-1 / slide-01 / Slide1).
    Settle on slide-001.png so the plan can address slides unambiguously."""
    out = []
    for p in sorted(out_dir.glob("*.png"),
                    key=lambda q: int(re.search(r"(\d+)", q.stem).group(1))):
        n = int(re.search(r"(\d+)", p.stem).group(1))
        target = out_dir / f"slide-{n:03d}.png"
        if p != target:
            p.replace(target)
        out.append(target)
    return out


def powerpoint_available() -> bool:
    if sys.platform != "win32":
        return False
    try:
        import win32com.client                                    # noqa: F401
    except ImportError:
        return False
    try:
        import win32com.client as w
        app = w.Dispatch("PowerPoint.Application")
        app.Quit()
        return True
    except Exception:                                             # noqa: BLE001
        return False


def render_pngs_powerpoint(pptx: Path, out_dir: Path, dpi: int) -> list[Path]:
    """Native PowerPoint export — the highest-fidelity option on Windows, and
    it needs neither LibreOffice nor poppler. Requires `pip install pywin32`."""
    import win32com.client
    import pythoncom

    pythoncom.CoInitialize()
    app = None
    pres = None
    try:
        app = win32com.client.Dispatch("PowerPoint.Application")
        # PowerPoint refuses Visible=False; opening WithWindow=False is enough
        # to keep it out of the way.
        pres = app.Presentations.Open(str(pptx.resolve()),
                                      ReadOnly=True, WithWindow=False)

        # PageSetup is in points (1/72 in) -> pixels at the requested DPI.
        width = int(round(pres.PageSetup.SlideWidth / 72.0 * dpi))
        height = int(round(pres.PageSetup.SlideHeight / 72.0 * dpi))

        for p in out_dir.glob("*.png"):
            p.unlink()
        pres.Export(str(out_dir.resolve()), "PNG", width, height)
    finally:
        try:
            if pres is not None:
                pres.Close()
        except Exception:                                         # noqa: BLE001
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:                                         # noqa: BLE001
            pass
        pythoncom.CoUninitialize()

    pngs = normalise_names(out_dir)
    if not pngs:
        raise RuntimeError("PowerPoint exported no images")
    return pngs


def render_pngs(pptx: Path, out_dir: Path, dpi: int) -> list[Path]:
    """pptx -> PDF (LibreOffice) -> PNG per page (pdftoppm)."""
    soffice = which_or_die("soffice", "libreoffice")
    pdftoppm = which_or_die("pdftoppm")

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        # LibreOffice is chatty and needs its own profile dir to run headless twice
        subprocess.run(
            [soffice, "--headless", "--norestore",
             f"-env:UserInstallation=file://{tmp_path / 'lo-profile'}",
             "--convert-to", "pdf", "--outdir", str(tmp_path), str(pptx)],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            timeout=600,
        )
        pdfs = list(tmp_path.glob("*.pdf"))
        if not pdfs:
            raise RuntimeError(f"LibreOffice produced no PDF for {pptx.name}")

        subprocess.run(
            [pdftoppm, "-png", "-r", str(dpi), str(pdfs[0]), str(out_dir / "slide")],
            check=True, timeout=900,
        )

    return normalise_names(out_dir)


# -------------------------------------------------------------- xml dump --

def shape_text(shape) -> list[dict]:
    """Paragraphs with indent level and bold/size hints — enough to infer
    heading vs body vs bullet without guessing from the image."""
    out = []
    if not shape.has_text_frame:
        return out
    for para in shape.text_frame.paragraphs:
        text = "".join(r.text for r in para.runs) or para.text or ""
        if not text.strip():
            continue
        sizes = [r.font.size.pt for r in para.runs if r.font.size]
        bolds = [bool(r.font.bold) for r in para.runs if r.font.bold is not None]
        out.append({
            "text": text.strip(),
            "level": para.level,
            "size_pt": max(sizes) if sizes else None,
            "bold": any(bolds) if bolds else None,
        })
    return out


def table_grid(shape) -> dict | None:
    if not getattr(shape, "has_table", False):
        return None
    t = shape.table
    rows = [[cell.text.strip() for cell in row.cells] for row in t.rows]
    return {
        "n_rows": len(rows),
        "n_cols": len(rows[0]) if rows else 0,
        "headers": rows[0] if rows else [],
        "rows": rows[1:] if len(rows) > 1 else [],
    }


def emu_box(shape) -> dict | None:
    try:
        return {
            "x": round(Emu(shape.left).inches, 2), "y": round(Emu(shape.top).inches, 2),
            "w": round(Emu(shape.width).inches, 2), "h": round(Emu(shape.height).inches, 2),
        }
    except (TypeError, ValueError, AttributeError):
        return None


def walk(shapes, media_dir: Path, counter: dict, acc: dict) -> None:
    """Recurse groups so nothing nested is missed."""
    for shape in shapes:
        if shape.shape_type == 6 and hasattr(shape, "shapes"):      # GROUP
            walk(shape.shapes, media_dir, counter, acc)
            continue

        entry = {
            "name": shape.name,
            "type": str(shape.shape_type),
            "box_in": emu_box(shape),
        }

        paras = shape_text(shape)
        if paras:
            entry["paragraphs"] = paras

        grid = table_grid(shape)
        if grid:
            entry["table"] = grid
            acc["tables"].append(grid)

        # PICTURE (13) — also covers picture placeholders.
        # NB: do NOT use hasattr(shape, "image") here. `image` is a property that
        # raises ValueError("no embedded image") for linked-but-not-embedded
        # pictures, and hasattr() only swallows AttributeError — so the guard
        # itself would throw. Detect by class/type, then access inside try.
        if shape.__class__.__name__ in ("Picture", "PlaceholderPicture"):
            try:
                img = shape.image
                counter["img"] += 1
                ext = (img.ext or "png").lstrip(".")
                fn = f"img-{counter['img']:03d}.{ext}"
                (media_dir / fn).write_bytes(img.blob)
                entry["image_file"] = f"media/{fn}"
                acc["images"].append(fn)
            except (ValueError, AttributeError, KeyError) as exc:
                entry["image_missing"] = str(exc)   # linked image, not embedded
                acc["image_errors"] += 1

        if getattr(shape, "has_chart", False):
            entry["chart"] = True
            acc["charts"] += 1

        if (paras or grid or entry.get("chart")
                or "image_file" in entry or "image_missing" in entry):
            acc["shapes"].append(entry)


def dump_xml(pptx: Path, out_dir: Path) -> dict:
    prs = Presentation(str(pptx))
    media_dir = out_dir / "media"
    media_dir.mkdir(exist_ok=True)
    counter = {"img": 0}

    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        acc = {"shapes": [], "tables": [], "images": [], "charts": 0, "image_errors": 0}
        walk(slide.shapes, media_dir, counter, acc)

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()

        slides.append({
            "n": i,
            "layout": slide.slide_layout.name,
            "png": f"slide-{i:03d}.png",
            "shapes": acc["shapes"],
            "n_tables": len(acc["tables"]),
            "n_images": len(acc["images"]),
            "n_charts": acc["charts"],
            "notes": notes,
        })

    return {
        "slide_w_in": round(Emu(prs.slide_width).inches, 2),
        "slide_h_in": round(Emu(prs.slide_height).inches, 2),
        "slides": slides,
    }


# ------------------------------------------------------------------ main --

def render(pptx: Path, out_dir: Path, dpi: int, renderer: str) -> tuple[list[Path], str]:
    """renderer: auto | powerpoint | libreoffice"""
    if renderer == "powerpoint":
        return render_pngs_powerpoint(pptx, out_dir, dpi), "powerpoint"
    if renderer == "libreoffice":
        return render_pngs(pptx, out_dir, dpi), "libreoffice"

    # auto — prefer native PowerPoint on Windows, fall back to LibreOffice.
    if powerpoint_available():
        try:
            return render_pngs_powerpoint(pptx, out_dir, dpi), "powerpoint"
        except Exception as exc:                                  # noqa: BLE001
            print(f"  powerpoint export failed ({exc}); trying libreoffice",
                  flush=True)
    return render_pngs(pptx, out_dir, dpi), "libreoffice"


def extract_one(pptx: Path, out_root: Path, dpi: int, source_root: Path | None,
                resume: bool = False, renderer: str = "auto") -> dict:
    slug = slugify(pptx.name)
    out_dir = out_root / slug
    out_dir.mkdir(parents=True, exist_ok=True)

    # Parse first — it is fast, and its slide count is what tells us whether a
    # cached render is complete or was cut short.
    print("  parsing    xml", flush=True)
    data = dump_xml(pptx, out_dir)
    n_xml = len(data["slides"])

    cached = existing_pngs(out_dir)
    engine = "cache"
    if resume and len(cached) == n_xml and n_xml > 0:
        print(f"  reusing    {len(cached)} rendered pages", flush=True)
        pngs = normalise_names(out_dir)
    else:
        if resume and cached:
            print(f"  re-render  cache has {len(cached)} of {n_xml} pages "
                  f"— incomplete, discarding", flush=True)
            for p in cached:
                p.unlink()
        print(f"  rendering  {n_xml} slides", flush=True)
        pngs, engine = render(pptx, out_dir, dpi, renderer)
        print(f"  rendered   via {engine}", flush=True)

    n_png = len(pngs)
    if n_xml != n_png:
        # Usually a hidden slide (LibreOffice omits those from PDF export) or a
        # render that was cut short. Either way, never silently proceed.
        print(f"  !! slide count mismatch: xml={n_xml} png={n_png}", flush=True)

    (out_dir / "slides.json").write_text(
        json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    chapter = "?"
    if source_root:
        try:
            chapter = pptx.relative_to(source_root).parts[0]
        except ValueError:
            chapter = pptx.parent.name

    manifest = {
        "slug": slug,
        "source": str(pptx),
        "source_name": pptx.name,
        "chapter": chapter,
        "slide_count": n_xml,          # authoritative — the validator uses this
        "png_count": n_png,
        "dpi": dpi,
        "renderer": engine,
        "slide_w_in": data["slide_w_in"],
        "slide_h_in": data["slide_h_in"],
        "has_tables": sum(s["n_tables"] for s in data["slides"]),
        "has_images": sum(s["n_images"] for s in data["slides"]),
        "has_charts": sum(s["n_charts"] for s in data["slides"]),
        "count_ok": n_xml == n_png,
    }
    (out_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"  done       {slug}  ({n_xml} slides, "
          f"{manifest['has_tables']} tables, {manifest['has_images']} images)",
          flush=True)
    return manifest


def main() -> None:
    ap = argparse.ArgumentParser(description="Extract PPTX slides to PNG + JSON.")
    ap.add_argument("path", help=".pptx file, or a folder when using --all")
    ap.add_argument("--out", default="build", help="output root (default: build)")
    ap.add_argument("--dpi", type=int, default=150, help="render DPI (default: 150)")
    ap.add_argument("--all", action="store_true", help="recurse the folder for *.pptx")
    ap.add_argument("--resume", action="store_true",
                    help="reuse already-rendered PNGs (skip the slow render step)")
    ap.add_argument("--renderer", choices=("auto", "powerpoint", "libreoffice"),
                    default="auto",
                    help="auto (default): native PowerPoint on Windows, else LibreOffice")
    args = ap.parse_args()

    src = Path(args.path)
    out_root = Path(args.out)
    out_root.mkdir(parents=True, exist_ok=True)

    if args.all:
        if not src.is_dir():
            sys.exit(f"--all needs a folder, got {src}")
        targets = sorted(p for p in src.rglob("*.pptx") if not is_lock_file(p))
        source_root = src
    else:
        if not src.is_file():
            sys.exit(f"not a file: {src}")
        if is_lock_file(src):
            sys.exit(f"{src.name} is an Office lock file, not a deck")
        targets, source_root = [src], src.parent.parent

    if not targets:
        sys.exit("no .pptx files found")

    print(f"{len(targets)} deck(s) -> {out_root}/\n")
    results, failures = [], []
    for i, t in enumerate(targets, 1):
        print(f"[{i}/{len(targets)}] {t.name}")
        try:
            results.append(extract_one(t, out_root, args.dpi, source_root,
                                       args.resume, args.renderer))
        except Exception as exc:                                   # noqa: BLE001
            print(f"  FAILED: {exc}", flush=True)
            failures.append({"source": str(t), "error": str(exc)})

    status = out_root / "status.csv"
    write_header = not status.exists()
    with status.open("a", encoding="utf-8") as fh:
        if write_header:
            fh.write("slug,chapter,slides,tables,images,count_ok,phase\n")
        for m in results:
            fh.write(f"{m['slug']},{m['chapter']},{m['slide_count']},"
                     f"{m['has_tables']},{m['has_images']},{m['count_ok']},extracted\n")

    print(f"\nextracted {len(results)}, failed {len(failures)}")
    for f in failures:
        print(f"  FAIL {f['source']}: {f['error']}")
    if failures:
        sys.exit(1)


if __name__ == "__main__":
    main()
